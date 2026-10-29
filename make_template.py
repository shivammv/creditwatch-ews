import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as L

RAW = ["total_assets","current_assets","current_liabilities","total_liabilities","equity","retained_earnings","sales","cogs","sga",
       "depreciation","ebit","interest_expense","pbt","net_income","cfo","receivables","inventory","ppe_net","lt_debt","cash","shares_out"]
tmpl = pd.read_csv("data/indian/financials_TEMPLATE.csv")
ev = pd.read_csv("data/indian/events.csv"); pe = pd.read_csv("data/indian/peers.csv")
F = "Arial"
hdr = PatternFill("solid", fgColor="1F3864"); inp = PatternFill("solid", fgColor="FFF9DB"); ex = PatternFill("solid", fgColor="E2EFDA")
thin = Side(style="thin", color="BFBFBF"); bd = Border(left=thin, right=thin, top=thin, bottom=thin)
blue = Font(name=F, size=10, color="0000FF"); blk = Font(name=F, size=10); wh = Font(name=F, size=10, bold=True, color="FFFFFF")

wb = Workbook()
# ---------------- Instructions
ws = wb.active; ws.title = "Instructions"
lines = [
 ("CreditWatch EWS - Input Workbook", True),
 ("Fill only the YELLOW cells with BLUE text. Green row = worked example (ignored by the loader, names starting 'EXAMPLE' are skipped).", False),
 ("", False),
 ("SHEETS", True),
 ("Financials  - one row per company per fiscal year (3 years before event + 1 earlier year for Beneish/Piotroski). Pre-filled with the 10 defaulters and 32 matched peers.", False),
 ("Events      - documented event date per defaulter. Set date_status to 'verified' only after checking IBBI / NCLT / RBI / exchange record.", False),
 ("Peers       - 32 same-sector survivors matched to a defaulter. Remove/swap any peer that defaulted before the matched event date.", False),
 ("SoftSignals - one row per observation: rating_downgrade, promoter_pledge_high (value = pledge %), auditor_red_flag, delayed_filing, related_party_stress.", False),
 ("Ratio_Check - LIVE formulas on Financials: balance-sheet check, Altman Z'', current ratio, interest cover. Red 'CHECK' = likely typo or wrong unit.", False),
 ("", False),
 ("RULES", True),
 ("1. One currency unit for the whole file (INR crore recommended). Do not mix lakh/crore/million.", False),
 ("2. public_date = the date results / annual report became PUBLIC (leakage control). Blank => fy_end + 120 days.", False),
 ("3. Use only what was public BEFORE the event. Never paste post-event restated numbers.", False),
 ("4. Banks / NBFCs: put total income in 'sales'; Z'', M-score, F-score are skipped for them by design.", False),
 ("5. Unknown item? Leave blank. The pipeline derives total_liabilities, ebit, pbt where it can and handles other gaps natively.", False),
 ("6. Then run:  python -c \"from src.ingest import export_template as e; print(e('CreditWatch_Input_Template.xlsx'))\"  followed by  python run_all.py", False),
 ("   OR score any book directly:  python score_any.py <yourfile>.xlsx --asof 2026-10-27", False),
 ("", False),
 ("Colour legend: yellow fill + blue text = input | green = example | dark blue = headers | black = formula", False),
]
for i, (t, b) in enumerate(lines, 1):
    c = ws.cell(i, 1, t); c.font = Font(name=F, size=14 if i == 1 else 10, bold=b); c.alignment = Alignment(wrap_text=True, vertical="top")
ws.column_dimensions["A"].width = 150

def sheet(name, cols, rows, widths=None, example=None, input_from=0, extra_blank=0):
    s = wb.create_sheet(name)
    for j, c in enumerate(cols, 1):
        h = s.cell(1, j, c); h.font = wh; h.fill = hdr; h.alignment = Alignment(horizontal="center", wrap_text=True); h.border = bd
        s.column_dimensions[L(j)].width = (widths or {}).get(c, 16)
    r = 2
    if example:
        for j, v in enumerate(example, 1):
            c = s.cell(r, j, v); c.fill = ex; c.font = blk; c.border = bd
        r += 1
    for row in rows:
        for j, v in enumerate(row, 1):
            c = s.cell(r, j, None if pd.isna(v) else v); c.border = bd
            if j > input_from: c.fill = inp; c.font = blue
            else: c.font = blk
        r += 1
    for _ in range(extra_blank):
        for j in range(1, len(cols) + 1):
            c = s.cell(r, j); c.border = bd; c.fill = inp; c.font = blue
        r += 1
    s.freeze_panes = "B2"
    return s, r

# ---------------- Financials
fcols = ["company","role","matched_to","fy_end","public_date","source_url"] + RAW
exrow = ["EXAMPLE Ltd","defaulter","EXAMPLE Ltd","2018-03-31","2018-08-10","https://www.bseindia.com/... (annual report link)",
         1000,400,300,600,400,200,1200,800,100,50,150,40,110,80,120,150,100,450,250,60,100]
rows = tmpl[fcols].values.tolist()
fs, endr = sheet("Financials", fcols, rows, {"company":26,"source_url":34,"matched_to":22}, exrow, input_from=3)
for r in range(3, endr):
    fs.cell(r, 4).number_format = "yyyy-mm-dd"; fs.cell(r, 5).number_format = "yyyy-mm-dd"
    for j in range(7, 28): fs.cell(r, j).number_format = "#,##0.0;(#,##0.0);-"
dv = DataValidation(type="list", formula1='"defaulter,peer"', allow_blank=True); fs.add_data_validation(dv); dv.add(f"B3:B{endr}")
nf = endr - 1

# ---------------- Events
ecols = list(ev.columns)
es, er = sheet("Events", ecols, ev.values.tolist(), {"company":26,"source_to_check":42,"note":60,"event_type":34}, None, input_from=3)
dv2 = DataValidation(type="list", formula1='"unverified,verified"', allow_blank=True); es.add_data_validation(dv2)
dv2.add(f"{L(ecols.index('date_status')+1)}2:{L(ecols.index('date_status')+1)}{er}")

# ---------------- Peers
ps, _ = sheet("Peers", list(pe.columns), pe.values.tolist(), {"peer":28,"status_check":58}, None, input_from=3)

# ---------------- SoftSignals
ss, sr = sheet("SoftSignals", ["company","signal","date_public","value","source_url"], [], {"company":26,"signal":24,"source_url":40},
               ["EXAMPLE Ltd","rating_downgrade","2018-05-14",1,"https://www.crisil.com/... (rationale link)"], extra_blank=60)
dv3 = DataValidation(type="list", formula1='"rating_downgrade,promoter_pledge_high,auditor_red_flag,delayed_filing,related_party_stress"', allow_blank=True)
ss.add_data_validation(dv3); dv3.add("B3:B200")

# ---------------- Ratio_Check (live formulas)
rc = wb.create_sheet("Ratio_Check")
rcols = ["company","fy_end","Balance gap (TA-TL-EQ)","Gap % of TA","Working capital / TA","Retained / TA","EBIT / TA","Equity / TL","Altman Z'' (non-financial)","Zone","Current ratio","Interest cover (EBIT/Int)","Data check"]
for j, c in enumerate(rcols, 1):
    h = rc.cell(1, j, c); h.font = wh; h.fill = hdr; h.alignment = Alignment(horizontal="center", wrap_text=True); rc.column_dimensions[L(j)].width = 17
rc.column_dimensions["A"].width = 26; rc.row_dimensions[1].height = 42
# Financials cols: G=TA H=CA I=CL J=TL K=EQ L=RE M=Sales N=COGS O=SGA P=Dep Q=EBIT R=Int
for r in range(2, endr):
    f = lambda col: f"Financials!{col}{r}"
    vals = [f"=IF({f('A')}=\"\",\"\",{f('A')})", f"=IF({f('D')}=\"\",\"\",{f('D')})",
            f"=IF(COUNT({f('G')},{f('J')},{f('K')})<3,\"\",{f('G')}-{f('J')}-{f('K')})",
            f"=IF(C{r}=\"\",\"\",IFERROR(C{r}/{f('G')},\"\"))",
            f"=IF(COUNT({f('H')},{f('I')},{f('G')})<3,\"\",IFERROR(({f('H')}-{f('I')})/{f('G')},\"\"))",
            f"=IF(COUNT({f('L')},{f('G')})<2,\"\",IFERROR({f('L')}/{f('G')},\"\"))",
            f"=IF(COUNT({f('Q')},{f('G')})<2,\"\",IFERROR({f('Q')}/{f('G')},\"\"))",
            f"=IF(COUNT({f('K')},{f('J')})<2,\"\",IFERROR({f('K')}/{f('J')},\"\"))",
            f"=IF(COUNT(E{r}:H{r})<4,\"\",6.56*E{r}+3.26*F{r}+6.72*G{r}+1.05*H{r})",
            f"=IF(I{r}=\"\",\"\",IF(I{r}>2.6,\"Safe\",IF(I{r}>=1.1,\"Grey\",\"Distress\")))",
            f"=IF(COUNT({f('H')},{f('I')})<2,\"\",IFERROR({f('H')}/{f('I')},\"\"))",
            f"=IF(COUNT({f('Q')},{f('R')})<2,\"\",IFERROR({f('Q')}/{f('R')},\"\"))",
            f"=IF(C{r}=\"\",\"\",IF(ABS(D{r})>0.05,\"CHECK: assets <> liabilities+equity\",\"OK\"))"]
    for j, v in enumerate(vals, 1):
        c = rc.cell(r, j, v); c.font = blk
        if j in (3,): c.number_format = "#,##0.0;(#,##0.0);-"
        if j in (4,5,6,7): c.number_format = "0.0%"
        if j in (8,9,11,12): c.number_format = "0.00"
        if j == 2: c.number_format = "yyyy-mm-dd"
from openpyxl.formatting.rule import CellIsRule, FormulaRule
rc.conditional_formatting.add(f"M2:M{endr}", FormulaRule(formula=[f'LEFT(M2,5)="CHECK"'], font=Font(name=F, color="C00000", bold=True)))
rc.conditional_formatting.add(f"J2:J{endr}", FormulaRule(formula=['J2="Distress"'], font=Font(name=F, color="C00000", bold=True)))
rc.freeze_panes = "C2"
wb.save("/mnt/user-data/outputs/CreditWatch_Input_Template.xlsx")
print("rows", endr)
